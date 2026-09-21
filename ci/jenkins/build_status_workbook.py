#!/usr/bin/env python3
"""Generate the per-build ACT status workbook archived by Jenkins."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

STATUS_ORDER = ("PASS", "FAIL", "NOT RUN")
SHEET_NAMES = {
    "Privileged": "Privileged Tests",
    "Non-Privileged": "Non-Privileged Tests",
    "Vector": "Vector Tests",
}
HEADERS = (
    "Test Name",
    "Extension",
    "Suite",
    "Sail",
    "Spike",
    "Hardware",
    "Failure Category",
    "Failure Reason",
)
NAVY = "13253A"
BLUE = "235789"
CYAN = "18A999"
WHITE = "FFFFFF"
PALE_BLUE = "EAF2F8"
PASS_FILL = "D9EAD3"
FAIL_FILL = "F4CCCC"
NOT_RUN_FILL = "FFF2CC"
GRID = Side(style="thin", color="D7DEE8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--run-root", type=Path, required=True)
    parser.add_argument("--artifact-root", type=Path, required=True)
    parser.add_argument("--platform-label", required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--test-scope", choices=("priv", "all"), required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def normalized_status(value: object) -> str:
    status = str(value or "").strip().upper().replace("_", " ")
    if status in {"PASS", "PASSED", "SUCCESS"}:
        return "PASS"
    if status in {"FAIL", "FAILED", "FAILURE", "ERROR", "TIMEOUT"}:
        return "FAIL"
    return "NOT RUN"


def read_status_tsv(path: Path) -> dict[str, str]:
    statuses: dict[str, str] = {}
    if not path.is_file():
        return statuses
    with path.open(newline="", encoding="utf-8", errors="replace") as stream:
        for row in csv.reader(stream, delimiter="\t"):
            if len(row) >= 2 and row[0] and row[0] != "test_name":
                statuses[row[0]] = normalized_status(row[1])
    return statuses


def read_cases(path: Path) -> dict[str, dict]:
    if not path.is_file():
        return {}
    loaded = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(loaded, list):
        raise TypeError(f"{path} must contain a JSON list")
    return {
        str(case["test_name"]): case
        for case in loaded
        if isinstance(case, dict) and case.get("test_name")
    }


def test_name_from_artifact(path: Path) -> str:
    name = path.name
    for suffix in (".sig.elf", ".sig.log", ".sig", ".results", ".elf"):
        if name.endswith(suffix):
            return name[: -len(suffix)]
    return ""


def discover_test_metadata(
    artifact_root: Path, test_scope: str
) -> dict[str, tuple[str, str]]:
    metadata: dict[str, tuple[str, str]] = {}
    if not artifact_root.is_dir():
        return metadata
    for path in artifact_root.rglob("*"):
        if not path.is_file():
            continue
        name = test_name_from_artifact(path)
        if not name or name in metadata:
            continue
        relative = path.relative_to(artifact_root)
        lowered_parts = [part.lower() for part in relative.parts]
        extension = path.parent.name or re.sub(r"-\d+$", "", name)
        if test_scope == "priv" or "priv" in lowered_parts:
            suite = "Privileged"
        elif any(
            part in {"vector", "rv32v", "rv64v"} or part.startswith(("rv32v_", "rv64v_"))
            for part in lowered_parts
        ):
            suite = "Vector"
        else:
            suite = "Non-Privileged"
        metadata[name] = (extension, suite)
    return metadata


def status_counts(rows: list[dict], field: str) -> Counter:
    return Counter(row[field] for row in rows)


def pass_rate(counts: Counter) -> str:
    executed = counts["PASS"] + counts["FAIL"]
    return f"{counts['PASS'] * 100 / executed:.1f}%" if executed else "0.0%"


def style_status_cell(cell) -> None:
    fills = {
        "PASS": PatternFill("solid", fgColor=PASS_FILL),
        "FAIL": PatternFill("solid", fgColor=FAIL_FILL),
        "NOT RUN": PatternFill("solid", fgColor=NOT_RUN_FILL),
    }
    cell.fill = fills.get(str(cell.value), PatternFill())
    cell.alignment = Alignment(horizontal="center")
    cell.font = Font(bold=True)


def write_table_sheet(workbook: Workbook, title: str, rows: list[dict]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(tuple(row[key] for key in (
            "name", "extension", "suite", "sail", "spike", "hardware",
            "failure_category", "failure_reason",
        )))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{max(1, sheet.max_row)}"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=GRID)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 7)
        for column in (4, 5, 6):
            style_status_cell(row[column - 1])
    widths = (34, 24, 18, 13, 13, 13, 30, 72)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False


def write_dashboard(
    workbook: Workbook,
    rows: list[dict],
    platform_label: str,
    run_id: str,
) -> None:
    sheet = workbook.active
    sheet.title = "Dashboard"
    sheet.sheet_view.showGridLines = False
    for column, width in zip("ABCDEF", (28, 18, 18, 18, 18, 18)):
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:F1")
    sheet["A1"] = f"RISC-V ACT Results — {platform_label}"
    sheet["A1"].font = Font(color=WHITE, bold=True, size=20)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:F2")
    timestamp = datetime.now(timezone.utc).strftime("%b %d, %Y, %I:%M:%S %p UTC")
    sheet["A2"] = f"Report date and time: {timestamp}   •   Run: {run_id}"
    sheet["A2"].alignment = Alignment(horizontal="center")

    hardware = status_counts(rows, "hardware")
    executed = hardware["PASS"] + hardware["FAIL"]
    sections = [
        (5, f"{platform_label} — TEST STATUS"),
        (11, "REFERENCE MODEL STATUS"),
        (17, "RESULTS BY TEST GROUP"),
        (24, f"TOP {platform_label} FAILURE EXTENSIONS"),
    ]
    for row_number, label in sections:
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=6)
        cell = sheet.cell(row_number, 1, label)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)

    def write_row(row_number: int, values: list[object]) -> None:
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)

    write_row(6, ["Total tests", "Executed", "PASS", "FAIL", "NOT RUN", "Pass rate"])
    write_row(7, [
        len(rows), executed, hardware["PASS"], hardware["FAIL"], hardware["NOT RUN"],
        pass_rate(hardware),
    ])

    for cell in sheet[6]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)

    reference_rows = []
    for label, field in (("Sail", "sail"), ("Spike", "spike")):
        counts = status_counts(rows, field)
        reference_rows.append([
            label, len(rows), counts["PASS"], counts["FAIL"], counts["NOT RUN"],
            pass_rate(counts),
        ])
    write_row(12, ["Platform", "Total tests", "PASS", "FAIL", "NOT RUN", "Pass rate"])
    for row_number, values in enumerate(reference_rows, 13):
        write_row(row_number, values)

    write_row(18, ["Test group", "Total tests", "PASS", "FAIL", "NOT RUN", "Pass rate"])
    for row_number, suite in enumerate(SHEET_NAMES, 19):
        suite_rows = [row for row in rows if row["suite"] == suite]
        counts = status_counts(suite_rows, "hardware")
        write_row(row_number, [
            SHEET_NAMES[suite], len(suite_rows), counts["PASS"], counts["FAIL"],
            counts["NOT RUN"], pass_rate(counts),
        ])

    failure_extensions = Counter(
        row["extension"] for row in rows if row["hardware"] == "FAIL"
    )
    write_row(25, ["Extension", "Failures", "% of board failures"])
    total_failures = sum(failure_extensions.values())
    for row_number, (extension, failures) in enumerate(
        failure_extensions.most_common(10), 26
    ):
        write_row(row_number, [
            extension,
            failures,
            f"{failures * 100 / total_failures:.1f}%" if total_failures else "0.0%",
        ])

    for header_row in (6, 12, 18, 25):
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_rows(args: argparse.Namespace) -> list[dict]:
    sail = read_status_tsv(args.state_root / "sail_reference_status.tsv")
    spike = read_status_tsv(args.state_root / "spike_status.tsv")
    cases = read_cases(args.run_root / "cases.json")
    metadata = discover_test_metadata(args.artifact_root, args.test_scope)
    names = sorted(set(sail) | set(spike) | set(cases) | set(metadata), key=str.casefold)
    rows = []
    for name in names:
        case = cases.get(name, {})
        extension, suite = metadata.get(
            name,
            (re.sub(r"-\d+$", "", name), "Privileged" if args.test_scope == "priv" else "Non-Privileged"),
        )
        hardware = normalized_status(case.get("status"))
        failure_category = "" if hardware != "FAIL" else str(case.get("category", ""))
        failure_reason = "" if hardware != "FAIL" else str(case.get("root_cause", ""))
        rows.append(
            {
                "name": name,
                "extension": extension,
                "suite": suite,
                "sail": sail.get(name, "NOT RUN"),
                "spike": spike.get(name, "NOT RUN"),
                "hardware": hardware,
                "failure_category": failure_category,
                "failure_reason": failure_reason,
            }
        )
    return rows


def main() -> int:
    args = parse_args()
    args.state_root = args.state_root.resolve()
    args.run_root = args.run_root.resolve()
    args.artifact_root = args.artifact_root.resolve()
    rows = build_rows(args)
    workbook = Workbook()
    workbook.properties.creator = "Jenkins"
    workbook.properties.title = "ACT test status matrix"
    write_dashboard(workbook, rows, args.platform_label, args.run_id)
    write_table_sheet(workbook, "Test Status", rows)

    summary = workbook.create_sheet("Summary")
    summary.append(("Platform", "PASS", "FAIL", "NOT RUN"))
    for label, field in (("Sail", "sail"), ("Spike", "spike"), (args.platform_label, "hardware")):
        counts = status_counts(rows, field)
        summary.append((label, *(counts[status] for status in STATUS_ORDER)))
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
    summary.column_dimensions["A"].width = 28
    for column in "BCD":
        summary.column_dimensions[column].width = 14

    for suite, title in SHEET_NAMES.items():
        write_table_sheet(workbook, title, [row for row in rows if row["suite"] == suite])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"ACT status workbook: {args.output} ({len(rows)} tests)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
